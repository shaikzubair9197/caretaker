"""
Parses downloaded attachment files into structured data for DocumentViewerPopup.

PDF pages render lazily in the viewer (zoom-aware); all other formats are
eagerly parsed here so the viewer never touches the raw bytes directly.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path


# ── Format enum ───────────────────────────────────────────────────────────────

class DocFormat(Enum):
    DOCX = "docx"
    PDF = "pdf"
    PPTX = "pptx"
    XLSX = "xlsx"
    TXT = "txt"
    MD = "md"
    UNSUPPORTED = "unsupported"


# ── Parsed result types ───────────────────────────────────────────────────────

@dataclass
class DocxRun:
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    font_name: str = ""
    font_size: int = 0
    color: str = ""
    hyperlink: str = ""
    # Fidelity extensions (backward-compatible; renderers read these if present)
    highlight_color: str = ""


@dataclass
class DocxBlock:
    kind: str           # "heading" | "paragraph" | "bullet" | "table" | "image" | "header" | "footer"
    text: str = ""
    level: int = 0      # heading level (1-9) or bullet indent level
    alignment: str = "left"
    spacing_before: int = 0
    spacing_after: int = 0
    indent: int = 0
    runs: list[DocxRun] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)   # table only
    image_data: bytes | None = None
    image_width: int = 0
    image_height: int = 0
    hyperlink: str = ""
    # Fidelity extensions (backward-compatible; renderers read these if present)
    list_level: int = 0                    # explicit list nesting depth (0 = top-level)
    list_style: str = "bullet"             # "bullet" | "number"


@dataclass
class ParsedDocx:
    format: DocFormat = DocFormat.DOCX
    filename: str = ""
    blocks: list[DocxBlock] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    footers: list[str] = field(default_factory=list)


@dataclass
class ParsedPdf:
    """Holds only metadata; the viewer opens fitz.Document on demand."""
    format: DocFormat = DocFormat.PDF
    filename: str = ""
    file_path: Path = field(default_factory=Path)
    total_pages: int = 0


@dataclass
class PptxShape:
    shape_type: str
    left: int
    top: int
    width: int
    height: int
    text: str = ""
    runs: list[DocxRun] = field(default_factory=list)
    fill_color: str = ""
    line_color: str = ""
    image_data: bytes | None = None
    hyperlink: str = ""
    table: list[list[str]] | None = None


@dataclass
class PptxSlide:
    number: int
    title: str
    shapes: list[PptxShape] = field(default_factory=list)
    background_color: str = ""
    slide_width: int = 0
    slide_height: int = 0


@dataclass
class ParsedPptx:
    format: DocFormat = DocFormat.PPTX
    filename: str = ""
    slides: list[PptxSlide] = field(default_factory=list)


@dataclass
class XlsxCell:
    value: str = ""
    bold: bool = False
    italic: bool = False
    underline: bool = False
    font_name: str = ""
    font_size: int = 0
    font_color: str = ""
    fill_color: str = ""
    align: str = "left"
    valign: str = "top"
    colspan: int = 1
    rowspan: int = 1


@dataclass
class XlsxSheet:
    name: str
    columns: list[str]
    rows: list[list[XlsxCell]]
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)


@dataclass
class ParsedXlsx:
    format: DocFormat = DocFormat.XLSX
    filename: str = ""
    sheets: list[XlsxSheet] = field(default_factory=list)


@dataclass
class ParsedText:
    format: DocFormat = DocFormat.TXT
    filename: str = ""
    content: str = ""


@dataclass
class ParsedUnsupported:
    format: DocFormat = DocFormat.UNSUPPORTED
    filename: str = ""
    reason: str = ""


ParsedDoc = (
    ParsedDocx | ParsedPdf | ParsedPptx | ParsedXlsx | ParsedText | ParsedUnsupported
)

# ── Extension / content-type dispatch maps ────────────────────────────────────

_EXT_FORMAT: dict[str, DocFormat] = {
    ".docx": DocFormat.DOCX,
    ".pdf": DocFormat.PDF,
    ".pptx": DocFormat.PPTX,
    ".xlsx": DocFormat.XLSX,
    ".txt": DocFormat.TXT,
    ".md": DocFormat.MD,
    ".markdown": DocFormat.MD,
}

_CT_FORMAT: dict[str, DocFormat] = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocFormat.DOCX,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": DocFormat.PPTX,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": DocFormat.XLSX,
    "application/pdf": DocFormat.PDF,
    "text/plain": DocFormat.TXT,
    "text/markdown": DocFormat.MD,
}

_XLSX_ROW_LIMIT = 1_000


# ── Public entry point ────────────────────────────────────────────────────────

def parse(path: Path, filename: str, content_type: str = "") -> ParsedDoc:
    """
    Dispatch to the right parser based on file extension, then content-type.
    Returns a ParsedUnsupported if format is unknown or a library is missing.
    """
    ext = Path(filename).suffix.lower()
    fmt = _EXT_FORMAT.get(ext) or _CT_FORMAT.get(content_type.split(";")[0].strip())
    if fmt is None:
        return ParsedUnsupported(
            filename=filename,
            reason=f"Unsupported format (ext={ext!r}, content-type={content_type!r})",
        )

    match fmt:
        case DocFormat.DOCX:
            return _parse_docx(path, filename)
        case DocFormat.PDF:
            return _parse_pdf(path, filename)
        case DocFormat.PPTX:
            return _parse_pptx(path, filename)
        case DocFormat.XLSX:
            return _parse_xlsx(path, filename)
        case DocFormat.TXT | DocFormat.MD:
            return _parse_text(path, filename, fmt)
        case _:
            return ParsedUnsupported(filename=filename, reason="Unhandled format")


# ── Format parsers ────────────────────────────────────────────────────────────

def _docx_color(color) -> str:
    if color is None:
        return ""
    rgb = getattr(color, "rgb", None)
    if rgb is None:
        return ""
    return f"#{str(rgb)}"


def _docx_run(run) -> DocxRun:
    font = run.font
    color = _docx_color(getattr(font, "color", None))
    font_name = getattr(font, "name", None) or ""
    font_size = 0
    try:
        font_size = int(font.size.pt) if font.size is not None else 0
    except Exception:
        font_size = 0
    hyperlink = ""
    if getattr(run, "hyperlink", None) is not None:
        hyperlink = getattr(run.hyperlink, "target", "") or ""
    return DocxRun(
        text=run.text or "",
        bold=bool(getattr(font, "bold", False)),
        italic=bool(getattr(font, "italic", False)),
        underline=bool(getattr(font, "underline", False)),
        strike=bool(getattr(font, "strike", False)),
        font_name=font_name,
        font_size=font_size,
        color=color,
        hyperlink=hyperlink,
    )


def _docx_alignment(para) -> str:
    align = getattr(para, "alignment", None)
    if align is None:
        return "left"
    mapping = {
        0: "left",
        1: "center",
        2: "right",
        3: "justify",
    }
    return mapping.get(int(align), "left")


def _parse_docx(path: Path, filename: str) -> ParsedDocx | ParsedUnsupported:
    try:
        import docx as _docx
    except ImportError:
        return ParsedUnsupported(filename=filename, reason="python-docx not installed")
    try:
        doc = _docx.Document(str(path))
    except Exception as exc:
        return ParsedUnsupported(filename=filename, reason=f"DOCX open error: {exc}")

    blocks: list[DocxBlock] = []
    headers: list[str] = []
    footers: list[str] = []

    for section in doc.sections:
        if section.header is not None:
            header_text = "\n".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
            if header_text:
                headers.append(header_text)
        if section.footer is not None:
            footer_text = "\n".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
            if footer_text:
                footers.append(footer_text)

    for header in headers:
        blocks.append(DocxBlock(kind="header", text=header))

    for para in doc.paragraphs:
        if not para.text.strip() and not para.runs:
            continue
        style_name = getattr(getattr(para, "style", None), "name", None)
        style = (style_name or "").lower()
        runs = [_docx_run(run) for run in para.runs if run.text]
        if style.startswith("heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            blocks.append(DocxBlock(
                kind="heading",
                text=para.text,
                level=level,
                alignment=_docx_alignment(para),
                spacing_before=int(para.paragraph_format.space_before.pt) if para.paragraph_format.space_before else 0,
                spacing_after=int(para.paragraph_format.space_after.pt) if para.paragraph_format.space_after else 0,
                runs=runs,
            ))
        elif "list" in style or "bullet" in style:
            ilvl = (para._p.pPr.numPr.ilvl.val if
                    para._p.pPr is not None and
                    para._p.pPr.numPr is not None and
                    para._p.pPr.numPr.ilvl is not None else 0)
            is_numbered = "list number" in style or "ordered" in style
            blocks.append(DocxBlock(
                kind="bullet",
                text=para.text,
                level=int(ilvl),
                alignment=_docx_alignment(para),
                runs=runs,
                list_level=int(ilvl),
                list_style="number" if is_numbered else "bullet",
            ))
        else:
            blocks.append(DocxBlock(
                kind="paragraph",
                text=para.text,
                alignment=_docx_alignment(para),
                spacing_before=int(para.paragraph_format.space_before.pt) if para.paragraph_format.space_before else 0,
                spacing_after=int(para.paragraph_format.space_after.pt) if para.paragraph_format.space_after else 0,
                runs=runs,
            ))

    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            blocks.append(DocxBlock(kind="table", rows=rows))

    for shape in doc.inline_shapes:
        image_data = None
        try:
            image_data = shape.image.blob
        except Exception:
            image_data = None
        if image_data:
            blocks.append(DocxBlock(
                kind="image",
                image_data=image_data,
                image_width=int(shape.width / 12700),
                image_height=int(shape.height / 12700),
            ))

    for footer in footers:
        blocks.append(DocxBlock(kind="footer", text=footer))

    return ParsedDocx(filename=filename, blocks=blocks, headers=headers, footers=footers)


def _parse_pdf(path: Path, filename: str) -> ParsedPdf | ParsedUnsupported:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return ParsedUnsupported(filename=filename, reason="PyMuPDF (fitz) not installed")
    try:
        doc = fitz.open(str(path))
        total = doc.page_count
        doc.close()
    except Exception as exc:
        return ParsedUnsupported(filename=filename, reason=f"PDF open error: {exc}")

    return ParsedPdf(filename=filename, file_path=path, total_pages=total)


def _pptx_color(color) -> str:
    if color is None:
        return ""
    rgb = getattr(color, "rgb", None)
    if rgb is None:
        return ""
    return f"#{str(rgb)}"


def _pptx_run(run) -> DocxRun:
    font = run.font
    color = _pptx_color(getattr(font, "color", None))
    font_name = getattr(font, "name", None) or ""
    font_size = 0
    try:
        font_size = int(font.size.pt) if font.size is not None else 0
    except Exception:
        font_size = 0
    return DocxRun(
        text=run.text or "",
        bold=bool(getattr(font, "bold", False)),
        italic=bool(getattr(font, "italic", False)),
        underline=bool(getattr(font, "underline", False)),
        strike=bool(getattr(font, "strike", False)),
        font_name=font_name,
        font_size=font_size,
        color=color,
        hyperlink="",
    )


def _pptx_shape(shape) -> PptxShape:
    left = int(shape.left / 12700)
    top = int(shape.top / 12700)
    width = int(shape.width / 12700)
    height = int(shape.height / 12700)
    fill_color = ""
    line_color = ""
    try:
        fill = shape.fill
        if getattr(fill, "fore_color", None) is not None:
            fill_color = _pptx_color(fill.fore_color)
    except Exception:
        fill_color = ""
    try:
        line = shape.line
        if getattr(line, "color", None) is not None:
            line_color = _pptx_color(line.color)
    except Exception:
        line_color = ""

    image_data = None
    if getattr(shape, "image", None) is not None:
        try:
            image_data = shape.image.blob
        except Exception:
            image_data = None

    text = ""
    runs: list[DocxRun] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text or ""
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.text:
                    runs.append(_pptx_run(run))
            if paragraph is not shape.text_frame.paragraphs[-1]:
                runs.append(DocxRun(text="\n"))

    table: list[list[str]] | None = None
    if getattr(shape, "has_table", False):
        table = []
        for row in shape.table.rows:
            table.append([cell.text for cell in row.cells])

    return PptxShape(
        shape_type=str(shape.shape_type),
        left=left,
        top=top,
        width=width,
        height=height,
        text=text,
        runs=runs,
        fill_color=fill_color,
        line_color=line_color,
        image_data=image_data,
        table=table,
    )


def _parse_pptx(path: Path, filename: str) -> ParsedPptx | ParsedUnsupported:
    try:
        from pptx import Presentation
        from pptx.enum.dml import MSO_FILL_TYPE
    except ImportError:
        return ParsedUnsupported(filename=filename, reason="python-pptx not installed")
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return ParsedUnsupported(filename=filename, reason=f"PPTX open error: {exc}")

    slides: list[PptxSlide] = []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()

        background_color = ""
        try:
            bg = slide.background.fill
            if getattr(bg, "type", None) == MSO_FILL_TYPE.SOLID:
                background_color = _pptx_color(getattr(bg.fore_color, "rgb", None))
        except Exception:
            background_color = ""

        shapes: list[PptxShape] = []
        for shape in slide.shapes:
            shapes.append(_pptx_shape(shape))

        slides.append(PptxSlide(
            number=i,
            title=title,
            shapes=shapes,
            background_color=background_color,
            slide_width=int(prs.slide_width / 12700),
            slide_height=int(prs.slide_height / 12700),
        ))

    return ParsedPptx(filename=filename, slides=slides)


def _xlsx_color(color) -> str:
    if color is None:
        return ""
    rgb = getattr(color, "rgb", None)
    if rgb is None:
        return ""
    hex_value = str(rgb)
    if len(hex_value) == 8:
        hex_value = hex_value[2:]
    return f"#{hex_value}"


def _parse_xlsx(path: Path, filename: str) -> ParsedXlsx | ParsedUnsupported:
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
    except ImportError:
        return ParsedUnsupported(filename=filename, reason="openpyxl not installed")
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except Exception as exc:
        return ParsedUnsupported(filename=filename, reason=f"XLSX open error: {exc}")

    sheets: list[XlsxSheet] = []
    for raw_name in wb.sheetnames:
        name = str(raw_name).strip() if raw_name else ""
        if not name:
            name = f"Sheet {len(sheets) + 1}"
        ws = wb[name]

        cells: list[list[XlsxCell]] = []
        max_col = 0
        for row_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
            if row_idx > _XLSX_ROW_LIMIT + 1:
                break
            cells_row: list[XlsxCell] = []
            for cell in row:
                value = cell.value
                text = "" if value is None else str(value)
                font = getattr(cell, "font", None)
                fill = getattr(cell, "fill", None)
                alignment = getattr(cell, "alignment", None)
                cells_row.append(XlsxCell(
                    value=text,
                    bold=bool(getattr(font, "bold", False)),
                    italic=bool(getattr(font, "italic", False)),
                    underline=bool(getattr(font, "underline", False)),
                    font_name=getattr(font, "name", "") or "",
                    font_size=int(font.sz) if getattr(font, "sz", None) else 0,
                    font_color=_xlsx_color(getattr(font, "color", None)),
                    fill_color=_xlsx_color(getattr(fill, "fgColor", None)) if fill is not None else "",
                    align=(alignment.horizontal or "left").lower() if alignment is not None else "left",
                    valign=(alignment.vertical or "top").lower() if alignment is not None else "top",
                ))
            max_col = max(max_col, len(cells_row))
            cells.append(cells_row)

        columns = [get_column_letter(i) for i in range(1, max_col + 1)]
        merges: list[tuple[int, int, int, int]] = []
        for merge in ws.merged_cells.ranges:
            merges.append((merge.min_row - 1, merge.min_col - 1, merge.max_row - 1, merge.max_col - 1))

        sheets.append(XlsxSheet(name=name, columns=columns, rows=cells, merges=merges))

    wb.close()
    return ParsedXlsx(filename=filename, sheets=sheets)


def _parse_text(path: Path, filename: str, fmt: DocFormat) -> ParsedText | ParsedUnsupported:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return ParsedUnsupported(filename=filename, reason=f"Read error: {exc}")
    return ParsedText(format=fmt, filename=filename, content=content)
