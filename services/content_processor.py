"""
content_processor — pure document processing (Document Retrieval plan — Phase 1).

DocumentProcessor turns an IndexedContentCandidate's bytes into a
ProcessedDocument: extracted plain text (written to the on-disk cache, never to
Postgres — design rule #1), deterministic keywords, deterministic entities
(content_entities), and an embedding (embeddings.embedder). It is PURE with
respect to the database — it touches only the filesystem cache and the embedder,
never a DB session — so it is trivially testable and reusable by future consumers.

It owns the lifecycle version stamps (design rule #5):
  • PARSER_VERSION   — bump to force re-extraction of text.
  • PIPELINE_VERSION — bump to force re-run of keyword/entity/normalisation.
  • EMBEDDING_MODEL  — bump (via env) to force re-embedding.
The catalog compares these against each row to decide needs_reindex.

The canonical extracted-text cache path lives here (text_cache_path) so there is
ONE source of truth; content_catalog re-exports it (catalog -> processor, never
the reverse, keeping the processor DB-free).
"""

from __future__ import annotations

import hashlib
import io
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from embeddings.embedder import encode
from services.content_entities import EntityRegistry
from services.content_providers import IndexedContentCandidate
from utils.config import settings
from utils.logger import get_logger
from utils.time_utils import utcnow

logger = get_logger("services.content_processor")

# ── Lifecycle versions (design rule #5) ───────────────────────────────────────
PARSER_VERSION = "1"
PIPELINE_VERSION = "1"
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "all-MiniLM-L6-v2")  # mirrors embeddings.embedder

# Embed only the head of very long documents — the MiniLM model truncates to a
# few hundred tokens anyway; this just bounds memory (design: "oversized text
# truncated before embedding"). Keyword/entity extraction still use the full text.
_EMBED_MAX_CHARS = 20_000
_MAX_KEYWORDS = 25

# Stable base for resolving the relative cache dir regardless of process cwd.
_BASE_DIR = Path(__file__).resolve().parents[1]

# Compact English stopword set for deterministic keyword extraction.
_STOPWORDS = {
    "the", "and", "for", "are", "but", "not", "you", "all", "any", "can", "had",
    "her", "was", "one", "our", "out", "day", "get", "has", "him", "his", "how",
    "its", "may", "new", "now", "old", "see", "two", "way", "who", "did", "yes",
    "this", "that", "with", "from", "they", "will", "your", "what", "when",
    "which", "their", "there", "would", "could", "should", "about", "into",
    "than", "then", "them", "these", "those", "have", "been", "were", "also",
    "such", "only", "some", "more", "most", "other", "over", "very", "just",
    "like", "make", "made", "each", "between", "because", "while", "where",
    "after", "before", "above", "below", "again", "here", "both", "during",
}


class ProcessingError(Exception):
    """Raised when a candidate cannot be processed (unsupported, corrupt,
    oversized, decode failure). The indexer maps this to index_status=ERROR with
    capped retry/backoff."""


@dataclass
class ProcessedDocument:
    """The pure result of processing one candidate — everything the catalog needs
    to write/update a row, with NO database identity. The extracted text lives in
    the cache file at text_cache_path (relative to CONTENT_TEXT_CACHE_DIR)."""

    content_hash: str
    text_cache_path: str                 # relative filename under the cache dir, e.g. "<hash>.txt"
    char_count: int
    keywords: list[str] = field(default_factory=list)
    entities: list[dict] = field(default_factory=list)
    embedding: Optional[list[float]] = None
    embedding_model: Optional[str] = None
    embedding_dimension: Optional[int] = None
    parser_version: str = PARSER_VERSION
    pipeline_version: str = PIPELINE_VERSION
    text_extracted_at: Optional[object] = None
    embedded_at: Optional[object] = None


# ── Cache path (single source of truth; catalog re-exports this) ──────────────
def cache_dir() -> Path:
    raw = settings.CONTENT_TEXT_CACHE_DIR
    p = Path(raw)
    return p if p.is_absolute() else (_BASE_DIR / p)


def text_cache_path(content_hash: str) -> Path:
    """Absolute path of the extracted-text cache file for a content hash."""
    return cache_dir() / f"{content_hash}.txt"


# ── Per-format plain-text extraction (heavy libs imported lazily) ─────────────
def _extract_text(data: bytes, ext: str) -> str:
    ext = (ext or "").lower()
    if ext in (".txt",):
        return data.decode("utf-8", errors="ignore")
    if ext in (".md",):
        return _extract_markdown(data)
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    if ext == ".pptx":
        return _extract_pptx(data)
    if ext == ".xlsx":
        return _extract_xlsx(data)
    raise ProcessingError(f"unsupported extension: {ext}")


def _extract_markdown(data: bytes) -> str:
    text = data.decode("utf-8", errors="ignore")
    try:
        import markdown as _markdown
        html = _markdown.markdown(text)
        # Strip tags — keep the readable text for snippets/keywords.
        return re.sub(r"<[^>]+>", " ", html)
    except Exception as e:  # noqa: BLE001 - fall back to raw markdown
        logger.debug(f"Markdown render failed, using raw text: {e}")
        return text


def _extract_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"PyMuPDF unavailable: {e}")
    try:
        parts: list[str] = []
        with fitz.open(stream=data, filetype="pdf") as doc:
            for page in doc:
                parts.append(page.get_text())
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"PDF extraction failed: {e}")


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"python-docx unavailable: {e}")
    try:
        document = Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs if p.text)
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"DOCX extraction failed: {e}")


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"python-pptx unavailable: {e}")
    try:
        prs = Presentation(io.BytesIO(data))
        parts: list[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in para.runs)
                        if line:
                            parts.append(line)
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"PPTX extraction failed: {e}")


def _extract_xlsx(data: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"openpyxl unavailable: {e}")
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(" ".join(cells))
        wb.close()
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        raise ProcessingError(f"XLSX extraction failed: {e}")


# ── Deterministic keyword extraction ──────────────────────────────────────────
def _extract_keywords(text: str, filename: str) -> list[str]:
    tokens = re.findall(r"[a-z][a-z0-9_+#.\-]{2,}", text.lower())
    counts: dict[str, int] = {}
    for tok in tokens:
        tok = tok.strip(".-_")
        if len(tok) < 3 or tok in _STOPWORDS or tok.isdigit():
            continue
        counts[tok] = counts.get(tok, 0) + 1
    # Highest frequency first, ties broken alphabetically for determinism.
    ranked = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
    keywords = [tok for tok, _ in ranked[:_MAX_KEYWORDS]]
    # Always include distinctive filename-stem tokens (strong retrieval signal).
    stem = Path(filename).stem.lower()
    for tok in re.findall(r"[a-z][a-z0-9]{2,}", stem):
        if tok not in _STOPWORDS and tok not in keywords:
            keywords.append(tok)
    return keywords


class DocumentProcessor:
    """Pure processor: candidate bytes -> ProcessedDocument. No DB session."""

    def __init__(self, max_file_mb: Optional[float] = None) -> None:
        self._max_bytes = int((max_file_mb if max_file_mb is not None
                               else settings.CONTENT_MAX_FILE_MB) * 1024 * 1024)

    def process(self, candidate: IndexedContentCandidate) -> ProcessedDocument:
        # Oversized files are rejected BEFORE reading bytes (cheap; uses stat size)
        # so a huge file never blocks the worker. The indexer caps the resulting
        # ERROR retries.
        if self._max_bytes and candidate.size_bytes and candidate.size_bytes > self._max_bytes:
            raise ProcessingError(
                f"file exceeds CONTENT_MAX_FILE_MB ({candidate.size_bytes} bytes > {self._max_bytes})"
            )

        data = candidate.read_bytes()
        content_hash = hashlib.sha256(data).hexdigest()

        text = _extract_text(data, candidate.extension)
        normalized = re.sub(r"\s+", " ", text).strip()
        extracted_at = utcnow()

        # Persist the extracted text to the cache (never to Postgres — rule #1).
        rel_cache = self._write_cache(content_hash, text)

        keywords: list[str] = []
        entities: list[dict] = []
        embedding: Optional[list[float]] = None
        embedding_model: Optional[str] = None
        embedding_dimension: Optional[int] = None
        embedded_at = None

        if normalized:
            keywords = _extract_keywords(normalized, candidate.filename)
            entities = [e.as_dict() for e in EntityRegistry.extract_all(text)]
            vector = encode(normalized[:_EMBED_MAX_CHARS])
            if vector:
                embedding = vector
                embedding_model = EMBEDDING_MODEL
                embedding_dimension = len(vector)
                embedded_at = utcnow()
            else:
                # Embeddings unavailable — still a valid INDEXED row that serves on
                # keyword/entity/filename/recency signals (retrieval fallback).
                logger.info(f"No embedding for {candidate.filename} — embeddings unavailable")
        else:
            logger.info(f"Empty extracted text for {candidate.filename} — indexing metadata only")

        return ProcessedDocument(
            content_hash=content_hash,
            text_cache_path=rel_cache,
            char_count=len(text),
            keywords=keywords,
            entities=entities,
            embedding=embedding,
            embedding_model=embedding_model,
            embedding_dimension=embedding_dimension,
            parser_version=PARSER_VERSION,
            pipeline_version=PIPELINE_VERSION,
            text_extracted_at=extracted_at,
            embedded_at=embedded_at,
        )

    def _write_cache(self, content_hash: str, text: str) -> str:
        path = text_cache_path(content_hash)
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(text, encoding="utf-8")
        except OSError as e:
            raise ProcessingError(f"failed to write text cache {path}: {e}")
        return path.name  # relative to CONTENT_TEXT_CACHE_DIR
